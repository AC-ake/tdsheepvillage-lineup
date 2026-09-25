# -*- coding: utf-8 -*-
"""
从《保卫羊村百科全书V1.82.xlsx》与《保卫羊村全地图布局参考 2026.7.pdf》
提取布阵工具所需的数据，生成 data.js / maps/*.png / refs/*.jpg。

数据来源说明：
  * 地图格子来自工作表「前线、防线地图」里的“空白地图”。
    作者用单元格边框画格子（虚线 = 狼可走的格子，实线 = 不可走），
    用单元格填充色当“颜色索引”，图例写在每张地图下方的“注：”行。
  * 塔造价来自「各种塔造价、攻击力」（前线）与「防线塔造价、经验」（防线）。

重新生成： python extract_data.py
"""

import json
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, deque
from io import BytesIO

import openpyxl
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROJ = os.path.dirname(ROOT)
XLSX = os.path.join(PROJ, "保卫羊村百科全书V1.82.xlsx")
PDF = os.path.join(PROJ, "保卫羊村全地图布局参考 2026.7.pdf")
MEDIA_DIR = os.path.join(ROOT, "_build", "media")
MAPS_DIR = os.path.join(ROOT, "maps")
REFS_DIR = os.path.join(ROOT, "refs")

DASH_STYLES = {
    "dashed", "dotted", "dashDot", "dashDotDot",
    "mediumDashed", "mediumDashDot", "mediumDashDotDot", "slantDashDot",
}
THIN_STYLES = {"thin", "hair"}
FRAME_STYLES = {"medium", "thick", "double"}

ARROWS = {"\u2192": "right", "\u2190": "left", "\u2191": "up", "\u2193": "down"}

T_BLOCKED = 0
T_OPEN = 1
T_ENTRANCE = 2
T_VILLAGE = 3
T_TELE_IN = 4
T_TELE_OUT = 5
T_TOWER_ONLY = 6
T_WALK_ONLY = 7
T_MARK = 8

TYPE_NAME = {
    str(T_BLOCKED): "不可通过",
    str(T_OPEN): "可通行 / 可建塔",
    str(T_ENTRANCE): "出狼口",
    str(T_VILLAGE): "羊村入口",
    str(T_TELE_IN): "传送入",
    str(T_TELE_OUT): "传送出",
    str(T_TOWER_ONLY): "只能建塔（狼不可走）",
    str(T_WALK_ONLY): "只能走狼（不可建塔）",
    str(T_MARK): "标记格",
}


def unzip_media():
    os.makedirs(MEDIA_DIR, exist_ok=True)
    with zipfile.ZipFile(XLSX) as z:
        for name in z.namelist():
            if name.startswith("xl/media/"):
                out = os.path.join(MEDIA_DIR, os.path.basename(name))
                if not os.path.exists(out):
                    with open(out, "wb") as f:
                        f.write(z.read(name))


def read_drawing_anchors():
    ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    rid = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    with zipfile.ZipFile(XLSX) as z:
        rels = z.read("xl/drawings/_rels/drawing3.xml.rels").decode("utf-8")
        rel_map = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        root = ET.fromstring(z.read("xl/drawings/drawing3.xml"))
    anchors = []
    for anchor in list(root):
        frm = anchor.find("xdr:from", ns)
        if frm is None:
            continue
        col = int(frm.find("xdr:col", ns).text) + 1
        row = int(frm.find("xdr:row", ns).text) + 1
        media = [rel_map.get(b.get(rid)) for b in anchor.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}blip")]
        media = [os.path.basename(m) for m in media if m]
        if media:
            anchors.append({"col": col, "row": row, "media": media})
    return anchors


def fill_key(cell):
    if cell.fill is None or not cell.fill.patternType:
        return None
    fg = cell.fill.fgColor
    if fg is None:
        return None
    if fg.type == "rgb":
        return "rgb:%s" % fg.rgb
    if fg.type == "theme":
        return "theme:%s:%s" % (fg.theme, round(fg.tint or 0, 3))
    if fg.type == "indexed":
        return "idx:%s" % fg.indexed
    return None


def read_cell(cell):
    b = cell.border
    styles = [b.left.style, b.right.style, b.top.style, b.bottom.style]
    value = cell.value
    text = None
    if value is not None and not str(value).startswith("="):
        text = str(value).strip()
    return {
        "dashed": any(s in DASH_STYLES for s in styles if s),
        "thin": any(s in THIN_STYLES for s in styles if s),
        "frame": any(s in FRAME_STYLES for s in styles if s),
        "fill": fill_key(cell),
        "text": text,
    }


def find_blocks(ws):
    mask = {}
    for row in ws.iter_rows():
        for cell in row:
            info = read_cell(cell)
            if info["dashed"] or info["thin"] or info["frame"] or info["fill"]:
                mask[(cell.row, cell.column)] = info
    seen = set()
    blocks = []
    for start in list(mask):
        if start in seen:
            continue
        queue = deque([start])
        seen.add(start)
        cells = []
        while queue:
            cur = queue.popleft()
            cells.append(cur)
            r, c = cur
            for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1)):
                nb = (r + dr, c + dc)
                if nb in mask and nb not in seen:
                    seen.add(nb)
                    queue.append(nb)
        if len(cells) < 20:
            continue
        rows = [r for r, _ in cells]
        cols = [c for _, c in cells]
        blocks.append({
            "top": min(rows), "bottom": max(rows),
            "left": min(cols), "right": max(cols),
        })
    return blocks


def classify_legend(text):
    buildable = ("可以布防" in text) and ("无法布防" not in text)
    walkable = ("可以走狼" in text) and ("无法走狼" not in text)
    if walkable and buildable:
        return T_OPEN
    if buildable:
        return T_TOWER_ONLY
    if walkable:
        return T_WALK_ONLY
    # 图例里少数几条没有写“可以布防 / 可以走狼”，按语义判断
    if "\u91d1\u5b57\u5854" in text:      # 金字塔：狼按箭头穿过，两波互不交叉
        return T_WALK_ONLY
    if "\u5f39\u7c27\u843d\u70b9" in text:  # ※弹簧落点：不影响该格原有性质
        return T_OPEN
    return T_BLOCKED


def marker_type(text):
    if "入口" in text:
        return T_ENTRANCE
    if "羊村" in text:
        return T_VILLAGE
    if "传送入" in text:
        return T_TELE_IN
    if "传送出" in text:
        return T_TELE_OUT
    # ※ 只是「弹簧落点」的符号标记，格子本身的性质由底色决定
    # （落在钢网桥上的 ※ 不能布防，落在普通地面上的 ※ 可以布防），所以这里不返回 T_MARK
    return None


"""图册里有些格子用的是 Office 主题色填充，这里给出近似的 RGB，用来给传送阵认颜色。"""
THEME_RGB = {0: "FFFFFF", 1: "000000", 2: "E7E6E6", 3: "44546A", 4: "4472C4",
             5: "ED7D31", 6: "A5A5A5", 7: "FFC000", 8: "5B9BD5", 9: "70AD47"}


def fill_hex(fill):
    fs = str(fill or "")
    if re.fullmatch(r"[0-9A-Fa-f]{8}", fs):
        return "#" + fs[2:]
    if re.fullmatch(r"[0-9A-Fa-f]{6}", fs):
        return "#" + fs
    m = re.fullmatch(r"theme:(\d+):([\d.]+)", fs)
    if m:
        return "#" + THEME_RGB.get(int(m.group(1)), "888888")
    return ""


def parse_map(ws, block):
    cells = {}
    for r in range(block["top"], block["bottom"] + 1):
        for c in range(block["left"], block["right"] + 1):
            cells[(r, c)] = read_cell(ws.cell(row=r, column=c))

    legend = {}
    marker_colors = {}
    legend_names = {}
    for (r, c), info in cells.items():
        text = info["text"]
        fill = info["fill"]
        if not text or not fill:
            continue
        if len(text) > 10 and ("\uff08" in text or "(" in text):
            legend.setdefault(fill, classify_legend(text))
            legend_names.setdefault(fill, re.split(r"[\uff08(]", text)[0].strip())
        else:
            kind = marker_type(text)
            if kind is not None:
                marker_colors.setdefault(fill, kind)

    def cell_type(r, c):
        info = cells.get((r, c), {})
        text = info.get("text") or ""
        kind = marker_type(text)
        if kind is not None:
            return kind
        fill = info.get("fill")
        if fill and fill in legend:
            return legend[fill]
        if fill and fill in marker_colors:
            return marker_colors[fill]
        if info.get("dashed"):
            return T_OPEN
        return T_BLOCKED

    dashed = set()
    terrain = set()
    marker_cells = {}
    for (r, c), info in cells.items():
        text = info["text"] or ""
        fill = info["fill"]
        if info["dashed"]:
            dashed.add((r, c))
        if len(text) <= 10:
            if fill and fill in legend:
                terrain.add((r, c))
            kind = marker_type(text)
            if kind is not None and fill:
                marker_cells[(r, c)] = kind

    content = set(dashed) | terrain
    for (r, c) in marker_cells:
        if any((r + dr, c + dc) in content for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1))):
            content.add((r, c))
    if not content:
        return None

    r0 = min(r for r, _ in content)
    r1 = max(r for r, _ in content)
    c0 = min(c for _, c in content)
    c1 = max(c for _, c in content)

    base = set(content)
    edges = []
    DIRV = {"right": (0, 1), "left": (0, -1), "up": (-1, 0), "down": (1, 0)}
    outside = []
    for (r, c), kind in marker_cells.items():
        if (r, c) in base:
            continue
        text = cells[(r, c)]["text"] or ""
        direction = next((v for k, v in ARROWS.items() if k in text), None)
        outside.append((r, c, kind, direction))

    # 网格外的出入口：把网格扩到能包含它（地图不一定是矩形），
    # 并且把它顺着箭头走进/走出经过的格子标成「只能走狼」的通道
    gap = {}
    for (r, c, _kind, direction) in outside:
        r0 = min(r0, r); r1 = max(r1, r); c0 = min(c0, c); c1 = max(c1, c)
    for (r, c, _kind, direction) in outside:
        if not direction:
            continue
        step = DIRV[direction]
        p = (r + step[0], c + step[1])
        guard = 0
        while (p not in base) and r0 <= p[0] <= r1 and c0 <= p[1] <= c1 and guard < 12:
            gap[p] = T_WALK_ONLY
            p = (p[0] + step[0], p[1] + step[1])
            guard += 1

    # 这张图里真正用到的特殊地形（按颜色索引对上名字），供工具的图例用
    used_fills = []
    for info in cells.values():
        f = info.get("fill")
        if f and f in legend and f not in used_fills and legend_names.get(f):
            used_fills.append(f)
    legend_pairs = [[legend_names[f], legend[f]] for f in used_fills]

    # 金字塔是有方向性的通道：横着那截只许从左往右走，竖着那截只许从下往上走
    pyramid = {}
    pyr = [pos for pos, info in cells.items()
           if info.get("fill") and legend_names.get(info["fill"]) == "\u91d1\u5b57\u5854"
           and len(info.get("text") or "") <= 10]
    pyr_set = set(pyr)
    if pyr_set:
        for (r, c) in pyr_set:
            dirs = []
            if (r, c - 1) in pyr_set or (r, c + 1) in pyr_set:
                dirs.append("right")
            if (r - 1, c) in pyr_set or (r + 1, c) in pyr_set:
                dirs.append("up")
            pyramid["%d_%d" % (r - r0, c - c0)] = dirs or ["right", "up"]

    labels = {}
    tele_fill = {}
    feat = {}          # 每个格子的特殊地形名字（按填充色对上图例），例如 弹簧 / 钢网桥 / 有洞的钢网桥
    marks = {}         # 格子里的符号标记，例如 ※ = 弹簧落点
    spring = {}        # 弹簧落点里「在钢网桥上」的那部分（弹簧会把狼弹到这些格子）
    grid = []
    for r in range(r0, r1 + 1):
        row = []
        for c in range(c0, c1 + 1):
            kind = gap.get((r, c), cell_type(r, c))
            row.append(kind)
            info = cells.get((r, c), {})
            text = info.get("text")
            fill = info.get("fill")
            pos = "%d_%d" % (r - r0, c - c0)
            if text and text != "\u203b" and len(text) <= 10:
                labels[pos] = text
            # 特殊地形：同一大类（比如「只能走狼」）里还要分钢网桥 / 弹簧 / 爆炸箱……
            if fill and legend_names.get(fill) and marker_type(text or "") is None:
                feat[pos] = legend_names[fill]
            if (text or "").strip() == "\u203b":
                marks[pos] = "\u203b"
                # 弹簧落点：落在钢网桥上的才算「正式的落点」（也不能布防）
                if feat.get(pos) in ("钢网桥", "有洞的钢网桥"):
                    spring[pos] = 1
            if kind in (T_TELE_IN, T_TELE_OUT) and fill:
                tele_fill[(r - r0, c - c0)] = fill
        grid.append("".join(str(x) for x in row))

    # 传送口按颜色配对：同色的传送入 → 传送出
    color_ids = {}
    tele = {}
    tele_color = {}
    for pos in sorted(tele_fill):
        fill = tele_fill[pos]
        if fill not in color_ids:
            color_ids[fill] = len(color_ids) + 1
        tele["%d_%d" % pos] = color_ids[fill]
        tele_color["%d_%d" % pos] = fill_hex(fill)

    notes = []
    for r in range(block["top"], block["bottom"] + 3):
        for c in range(block["left"], block["right"] + 5):
            text = read_cell(ws.cell(row=r, column=c))["text"]
            if not text or len(text) < 6:
                continue
            if text.startswith("\u6ce8"):
                continue
            if text.startswith("\uff08") or text.startswith("("):
                if notes:
                    notes[-1] += text
                continue
            if "\uff08" in text:
                notes.append(text)
    deduped = []
    for n in notes:
        n = n.strip()
        if not n or n in deduped:
            continue
        if "点击" in n or "返回" in n or "\u5730\u56fe\u540d" in n or "\u672c\u56fe" in n:
            continue
        deduped.append(n)

    return {
        "rows": r1 - r0 + 1,
        "cols": c1 - c0 + 1,
        "grid": grid,
        "labels": labels,
        "tele": tele,
        "teleColor": tele_color,
        "feat": feat,
        "marks": marks,
        "spring": spring,
        "legend": legend_pairs,
        "pyramid": pyramid,
        "notes": deduped,
        "edges": edges,
    }


"""传送带方向：图册地图里的箭头是画在图上的，表里没有数据，所以按图人工核对。
   值是 "left" / "up" / "right" / "down"；字符串表示这张图所有传送带同向，
   字典表示按格子（"行_列"）分别指定。"""
BELT_DIR = {
    # 克勒蒙村：四条传送带顺时针从上往下依次是 ← ↑ → ↓（形成一圈逆时针的循环）
    "克勒蒙村": {"2_6": "left", "2_7": "left",
                "4_10": "up", "5_10": "up",
                "7_6": "right", "7_7": "right",
                "4_3": "down", "5_3": "down"},
    "佛里特镇": "right",
    # 泰威尔镇：上面 2 格 ↑、下面 4 格 →
    "泰威尔镇": {"3_8": "up", "4_8": "up",
                "6_4": "right", "6_5": "right", "6_6": "right", "6_9": "right"},
    "香树镇": "left",
}

"""个别格子的地形和游戏里对不上时的修正表（键是「行_列」，从 0 开始；值是正确的地形编号）。
   远古冰川 (7,8)（从 1 开始数）图册上画成了障碍，实际不是障碍，按实测改回可通行。"""
CELL_FIX = {
    "远古冰川": {"6_7": 1},
}


TOWER_DEFS = [
    ("shaota", "哨塔", 2, 3),
    ("sandanta", "散弹塔", 4, 5),
    ("paota", "炮塔", 6, 7),
    ("bodongta", "波动塔", 8, 9),
    ("xiangqianta", "镶嵌塔", 10, None),
]


def read_table(wb, sheet, cols, exp_cols=None):
    ws = wb[sheet]
    out = {}
    for key, name, cost_c, atk_c in cols:
        cost, atk, exp = [], [], []
        for r in range(4, ws.max_row + 1):
            level = ws.cell(row=r, column=1).value
            c = ws.cell(row=r, column=cost_c).value
            if not isinstance(level, (int, float)) or not isinstance(c, (int, float)):
                break
            cost.append(round(float(c)))
            if atk_c:
                a = ws.cell(row=r, column=atk_c).value
                atk.append(round(float(a)) if isinstance(a, (int, float)) else None)
            if exp_cols and key in exp_cols:
                e = ws.cell(row=r, column=exp_cols[key]).value
                exp.append(round(float(e)) if isinstance(e, (int, float)) else None)
        out[name] = {"cost": cost, "atk": atk, "exp": exp}
    return out


GEM_COLORS = {
    1: ("红", "#e04040"),
    2: ("黄", "#e8cf35"),
    3: ("绿", "#3fbf5f"),
    4: ("蓝", "#3f7fe0"),
    5: ("紫", "#8a5fd0"),
    6: ("黑", "#3a3a3a"),
}
GEM_TIERS = {1: "碎片", 2: "晶体", 3: "宝石", 4: "精华", 5: "强化"}
TOWER_KEYS = ["shaota", "sandanta", "paota", "bodongta", "xiangqianta"]
BASE_RANGE = {"哨塔": 1800, "散弹塔": 2200, "炮塔": 4000, "镶嵌塔": 2000, "波动塔": None}
# 波动塔是十字贯穿攻击：上下左右各 5 格（图册没有数值，按游戏里的表现设定，可改）
CROSS_RANGE = {"波动塔": 5}

# 《保卫羊村全地图布局参考 2026.7》里 1~46 号地图与四个大图的顺序
PDF_ORDER = [
    "巴罗村", "比丘村", "卡西村", "呼噜噜村", "安蒂亚村", "菲洛克村", "凯奇镇",
    "艾伊尔村", "拉帕斯村", "沃夫沼泽", "麦恩矿山",
    "埃特纳村", "雷尼尔村", "拉卡维村", "坦克尔村", "莫诺雷村", "奥兹玛村", "西利村",
    "塔特村", "苏兰德城", "赤火地带", "烈焰熔炉", "银蛇之巅", "特洛伊岛", "钓鱼岛",
    "齿轮工厂", "虎克工厂", "斯蒂尔镇", "拉瓦锡镇", "佛里特镇", "斯贝斯镇", "泰威尔镇",
    "泰勒斯镇", "克勒蒙村", "鲁姆车间", "弗兰克镇", "生命森林",
    "防线", "香树镇", "远古冰川", "缠怨谷", "金沙桥", "楼兰城", "梨花冰场",
    "凌绽雪原", "枫威瀑布",
]
# 大本营防线：和四个大图并列，单独一组（图册里和 38 号「防线」是同一张图，画了两遍）
EXTRA_MAPS = ["大本营防线"]


def read_gems():
    """30 种宝石（6 色 × 5 档）：镶嵌塔攻击公式、各塔射程倍率、售价。

    数据来自游戏配置 sys_config.json：
      * 镶嵌塔攻击 = ROUND((a + b*L + c*L^2) * damageRate + damageAdd)
        damageRate = 参数1 * 技能等级 + 1，damageAdd = 参数0 * 技能等级
      * 射程倍率 = 1 + 参数1 * 技能等级（ChangeRange 技能）
    该表与《保卫羊村百科全书V1.82》「各种塔造价、攻击力」的镶嵌塔各列完全一致。
    """
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    if not os.path.exists(src):
        return []
    cfg = json.load(open(src, encoding="utf-8"))
    gem_def = cfg["properties"]["gem"]
    packages = cfg["skill_package"]
    skills = cfg["skill"]["towerSkill"]
    towers = cfg["building"]["tower"]
    tower_name = {k: towers[k]["name"] for k in TOWER_KEYS}
    tower_key = {towers[k]["name"]: k for k in TOWER_KEYS}

    def own_skill(name, kind):
        """塔自己技能包里的效果（例如炮塔的溅射 AoeAP）。"""
        t = towers.get(tower_key.get(name, ""), {})
        pid = t.get("skill")
        if not pid or pid not in packages:
            return None
        for sid in packages[pid]["skills"]:
            info = skills.get(sid)
            if info and info["kindId"] == kind:
                return info["params"]
        return None

    def find_skill(gem, tower_key, kind):
        entry = gem["sp"].get(tower_key)
        if not entry:
            return None, 0
        for sid in packages[entry["id"]]["skills"]:
            info = skills.get(sid)
            if info and info["kindId"] == kind:
                return info["params"], entry["lv"]
        return None, 0

    out = [{
        "id": "none", "name": "无宝石", "color": 0, "colorName": "无",
        "tier": 0, "tierName": "", "hex": "#b9c2c7", "price": 0,
        "xi": {"a": 9, "b": 2.8, "c": 0.4205, "rate": 1, "add": 0},
        "range": {tower_name[k]: 1 for k in TOWER_KEYS},
        "rangeMin": {tower_name[k]: 0 for k in TOWER_KEYS},
        "aoe": {tower_name[k]: {"radii": (own_skill(tower_name[k], "AoeAP") or [0, 0])[0] if own_skill(tower_name[k], "AoeAP") else 0,
                                 "rate": (own_skill(tower_name[k], "AoeAP") or [0, 0])[1] if own_skill(tower_name[k], "AoeAP") else 0}
                for k in TOWER_KEYS},
    }]
    for key in ["hong", "lv", "huang", "zi", "lan", "hei"]:
        for tier in range(1, 6):
            gem = gem_def["%s%d" % (key, tier)]
            color_no = gem["color"]
            color_name, hex_color = GEM_COLORS[color_no]
            xi = {"a": 0, "b": 0, "c": 0, "rate": 1, "add": 0}
            params, level = find_skill(gem, "xiangqianta", "ChangeDam")
            if params:
                xi = {
                    "add": params[0] * level,
                    "rate": params[1] * level + 1,
                    "a": params[2], "b": params[3], "c": params[4],
                }
            rates = {}
            blind = {}
            aoe = {}
            for tk in TOWER_KEYS:
                name = tower_name[tk]
                rate = 1.0
                rmin = 0
                rp, rl = find_skill(gem, tk, "ChangeRange")
                if rp:
                    rate = 1 + rp[1] * rl
                    if len(rp) > 3 and rp[3] and rp[3] > 0:
                        rmin = rp[3] / 1000.0      # 无法攻击近处（近处盲区，单位：格）
                rates[name] = round(rate, 4)
                blind[name] = round(rmin, 3)
                # 溅射：宝石的技能包会顶掉塔自己的（例如黑宝石强化炮塔的溅射）
                ap, _lv = find_skill(gem, tk, "AoeAP")
                if not ap:
                    ap = own_skill(name, "AoeAP")
                aoe[name] = {"radii": (ap[0] if ap else 0), "rate": (ap[1] if ap and len(ap) > 1 else 0)}
            out.append({
                "id": "%s%d" % (key, tier),
                "name": gem["name"],
                "color": color_no,
                "colorName": color_name,
                "tier": gem["gamLevel"],
                "tierName": GEM_TIERS[gem["gamLevel"]],
                "hex": hex_color,
                "price": (gem.get("price") or {}).get("gold", 0),
                "xi": xi,
                "range": rates,
                "rangeMin": blind,
                "aoe": aoe,
            })
    return out


def read_gem_effects(wb):
    """从「宝石功能」表提取减益数值，给战斗模拟用：
       蓝宝石 = 冰霜减速（%），绿宝石 = 毒性伤害（每秒伤害 / 攻击力倍率 + 持续），
       红宝石 = 燃烧（每秒叠加伤害）。"""
    ws = wb["宝石功能"]
    cols = {"哨塔": 3, "散弹塔": 4, "炮塔": 6, "波动塔": 7}
    groups = {"黑宝石": 3, "红宝石": 10, "绿宝石": 18, "蓝宝石": 26, "紫宝石": 34, "黄宝石": 41}
    tiers = [("碎片", 1), ("晶体", 2), ("宝石", 3), ("精华", 4), ("强化", 5)]
    out = {"slow": {}, "poison": {}, "burn": {}, "slowXq": {}}

    def text_of(r, c):
        v = ws.cell(row=r, column=c).value
        return v if isinstance(v, str) else ""

    def find(t, pat):
        m = re.search(pat, t)
        return float(m.group(1)) if m else None

    for gem, base in groups.items():
        # 每种宝石的「基础效果」占 1~2 行，档位行位置不一，直接按 B 列文字找
        tier_rows = {}
        for r in range(base + 1, base + 12):
            if ws.cell(row=r, column=1).value is not None:
                break          # 下一组宝石开始了
            b = ws.cell(row=r, column=2).value
            if isinstance(b, str) and b.strip() in ("碎片", "晶体", "宝石", "精华", "强化"):
                tier_rows[b.strip()] = r
        for tier, _off in tiers:
            if tier not in tier_rows:
                continue
            r = tier_rows[tier]
            for tower, col in cols.items():
                t = text_of(r, col)
                if not t:
                    continue
                if gem == "蓝宝石":
                    v = find(t, r"叠加(\d+)%冰霜减速")
                    if v is not None:
                        out["slow"].setdefault(tower, {})[tier] = v
                elif gem == "绿宝石":
                    rec = {}
                    pct = find(t, r"攻击力的(\d+)%")
                    mul = find(t, r"攻击力的(\d+(?:\.\d+)?)倍")
                    if pct is not None:
                        rec["mul"] = pct / 100
                    if mul is not None:
                        rec["mul"] = mul
                    flat = find(t, r"每秒(\d+)点")
                    if flat is not None:
                        rec["perSec"] = flat
                    dur = find(t, r"持续(\d+)秒")
                    if dur is not None:
                        rec["dur"] = dur
                    if rec:
                        out["poison"].setdefault(tower, {})[tier] = rec
                elif gem == "红宝石" and tower in ("哨塔", "散弹塔"):
                    # 炮塔、波动塔的红宝石现在是胆怯效果，不再有燃烧
                    v = find(t, r"每秒(\d+)点叠加燃烧伤害")
                    if v is not None:
                        out["burn"].setdefault(tower, {})[tier] = v

    # 镶嵌塔的冰霜减速：表里按等级给出（1~30 在 J~N，31~60 在 P~T）
    # 黄宝石的暴击（哨塔/散弹塔）与眩晕（炮塔/波动塔）
    for tier, r in [("碎片",42),("晶体",43),("宝石",44),("精华",45),("强化",46)]:
        for tower, col, key in [("哨塔",3,"crit"),("散弹塔",4,"crit"),("炮塔",6,"vertigo"),("波动塔",7,"vertigo")]:
            t = text_of(r, col)
            if not t:
                continue
            if key == "crit":
                ch = find(t, r"(\d+)%几率([\d.]+)倍暴击")
                m2 = re.search(r"(\d+)%几率([\d.]+)倍暴击", t)
                if m2:
                    out.setdefault("crit", {}).setdefault(tower, {})[tier] = {
                        "chance": float(m2.group(1))/100, "mul": float(m2.group(2)),
                    }
            else:
                m2 = re.search(r"(\d+)%几率眩晕(\d+)秒", t)
                if m2:
                    out.setdefault("vertigo", {}).setdefault(tower, {})[tier] = {
                        "chance": float(m2.group(1))/100, "dur": float(m2.group(2)),
                    }
    # 黑宝石连击（哨塔/波动塔）、镶嵌塔诅咒；红宝石胆怯、紫宝石击退（炮塔/波动塔）
    for tier, r in [("碎片",4),("晶体",5),("宝石",6),("精华",7),("强化",8)]:
        for tower, col in [("哨塔",3),("波动塔",7),("镶嵌塔",5),("炮塔",6)]:
            t = text_of(r, col)
            if not t:
                continue
            m2 = re.search(r"连续(\d+)次攻击", t)
            if m2:
                out.setdefault("combo", {}).setdefault(tower, {})[tier] = int(m2.group(1))
            m2 = re.search(r"诅咒能量=塔等级\*?([\d.]+)?", t)
            if m2:
                out.setdefault("cuss", {}).setdefault(tower, {})[tier] = float(m2.group(1) or 1)
    for tier, r in [("碎片",12),("晶体",13),("宝石",14),("精华",15),("强化",16)]:
        for tower, col in [("炮塔",6),("波动塔",7)]:
            t = text_of(r, col)
            m2 = re.search(r"(\d+)%概率施加胆怯([\d.]+)秒", t)
            if m2:
                out.setdefault("timid", {}).setdefault(tower, {})[tier] = {
                    "chance": float(m2.group(1))/100, "dur": float(m2.group(2)),
                }
    for tier, r in [("碎片",35),("晶体",36),("宝石",37),("精华",38),("强化",39)]:
        for tower, col in [("炮塔",6),("波动塔",7)]:
            t = text_of(r, col)
            m2 = re.search(r"(\d+)%几率超长距离击退", t)
            if m2:
                out.setdefault("knock", {}).setdefault(tower, {})[tier] = float(m2.group(1))/100
        t = text_of(r, 5)                      # 紫宝石镶嵌塔：闪电弹射
        m2 = re.search(r"闪电弹射(\d+)次，伤害递减(\d+)%", t)
        if m2:
            out.setdefault("chain", {})[tier] = {"n": int(m2.group(1)), "dec": float(m2.group(2))/100}

    # 镶嵌塔的冰霜减速：表里按等级给出（1~30 在 J~N，31~60 在 P~T）
    xq_cols = [(10, 1), (11, 2), (12, 3), (13, 4), (14, 5)]
    xq_cols2 = [(16, 1), (17, 2), (18, 3), (19, 4), (20, 5)]
    tier_name = {1: "碎片", 2: "晶体", 3: "宝石", 4: "精华", 5: "强化"}
    for r in range(50, 80):
        lv1 = ws.cell(row=r, column=9).value
        lv2 = ws.cell(row=r, column=15).value
        for col, tier in xq_cols:
            if isinstance(lv1, (int, float)):
                v = ws.cell(row=r, column=col).value
                if isinstance(v, (int, float)):
                    out["slowXq"].setdefault(int(lv1), {})[tier_name[tier]] = float(v)
        for col, tier in xq_cols2:
            if isinstance(lv2, (int, float)):
                v = ws.cell(row=r, column=col).value
                if isinstance(v, (int, float)):
                    out["slowXq"].setdefault(int(lv2), {})[tier_name[tier]] = float(v)
    return out


def read_wolf_skills():
    """每种狼的抗性 / 免疫 / 特殊能力（走狼预测用）。
    数值取自怪物技能表，抗性单位是百分比的小数（2 = 200%）。"""
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    if not os.path.exists(src):
        return {}
    cfg = json.load(open(src, encoding="utf-8"))
    ms = cfg["skill"]["monsterSkill"]

    def value(skid, lev):
        info = ms.get(skid)
        if not info:
            return None
        lv = info.get("levels")
        if isinstance(lv, list) and lv and isinstance(lv[0], list):
            i = min(max(0, int(lev or 1) - 1), len(lv) - 1)
            return lv[i][0] if lv[i] else None
        p = info.get("params")
        return p[0] if isinstance(p, list) and p else None

    RES = ("resistFrost", "resistLight", "resistPoison", "resistVertigo", "resistFire",
           "resistSilence", "resistCrit", "resistBeat")
    WEAK = ("weakFrost", "weakLight", "weakPoison", "weakVertigo", "weakFire",
            "weakSilence", "weakCrit", "weakBeat")
    out = {}
    for wid, w in cfg["wolfs"].items():
        rec = {}
        for s in (w.get("skills") or []):
            info = ms.get(s.get("skid"))
            if not info:
                continue
            kind = info["kindId"]
            v = value(s.get("skid"), s.get("lev", 1))
            if kind in RES:
                key = kind[6:].lower()
                rec[key] = max(rec.get(key, 0), v or 0)
            elif kind in WEAK:
                key = kind[4:].lower()
                rec[key] = min(rec.get(key, 0), v or 0)
            elif kind == "fly":
                rec["fly"] = 1
            elif kind == "isolation":
                rec["light"] = max(rec.get("light", 0), 1)   # 技能「免疫电」：完全无视闪电攻击
            elif kind == "sprint":
                lv = info.get("levels") or []
                i = min(max(0, int(s.get("lev", 1)) - 1), len(lv) - 1) if lv else -1
                rec["sprint"] = (lv[i][-1] if i >= 0 and lv[i] else 0)
            elif kind in ("invisible", "shield", "reborn", "revive", "summon", "blink", "divide"):
                rec[kind] = 1
        if rec:
            out[wid] = rec
    return out


def read_wolf_skill_list():
    """每只狼的技能清单（带等级参数），给走狼模拟的技能时钟用。

    数据来自游戏配置的怪物技能表：每个技能有 kindId（技能种类）与 levels（各等级的参数行）。
    返回 {狼id: [{"k": 种类, "n": 名字, "lv": 等级, "p": 该等级的参数行}, ...]}
    抗性 / 怕性两类不重复放进这里（它们已经在 wolfSkill 里合并好了）。"""
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    if not os.path.exists(src):
        return {}
    cfg = json.load(open(src, encoding="utf-8"))
    ms = cfg["skill"]["monsterSkill"]
    skip = ("resist", "weak")
    out = {}
    for wid, w in cfg["wolfs"].items():
        lst = []
        for s in (w.get("skills") or []):
            info = ms.get(s.get("skid"))
            if not info:
                continue
            kind = info.get("kindId") or ""
            if kind.startswith(skip):
                continue
            levels = info.get("levels") or []
            lv = max(1, int(s.get("lev", 1) or 1))
            row = None
            if levels and isinstance(levels[0], list):
                row = levels[min(lv, len(levels)) - 1]
            lst.append({"k": kind, "n": info.get("name") or "", "lv": lv, "p": row})
        if lst:
            out[wid] = lst
    return out


def read_upgrade_time(wb):
    ws = wb["塔升级时间"]
    out = {}
    for name, col in [("哨塔", 2), ("散弹塔", 3), ("炮塔", 4), ("波动塔", 5), ("镶嵌塔", 6)]:
        arr = []
        for r in range(2, 2 + 130):
            v = ws.cell(row=r, column=col).value
            arr.append(round(float(v)) if isinstance(v, (int, float)) else None)
        out[name] = arr
    return out


def read_special(wb):
    front = wb["各种塔造价、攻击力"]
    defend = wb["防线塔造价、经验"]
    spec = [
        ("盘龙柱", 22, 12, 2, "上下左右各两格内的塔 +15% 攻速"),
        ("牛郎雕像", 23, 13, 1, "上下左右各两格内的塔 +20% 攻击力"),
        ("织女雕像", 24, 14, 1, "上下左右各两格内的塔 +20% 射程"),
        ("元宵灯楼", 25, 15, 9, "无特殊效果"),
        ("墙", 26, None, 99, "无特殊效果，主要用来挡路"),
    ]
    out = []
    for name, fc, dc, limit, effect in spec:
        f = front.cell(row=4, column=fc).value
        d = defend.cell(row=4, column=dc).value if dc else None
        out.append({
            "name": name,
            "front": round(float(f)) if isinstance(f, (int, float)) else None,
            "defend": round(float(d)) if isinstance(d, (int, float)) else None,
            "limit": limit,
            "effect": effect,
        })
    return out


def read_tower_max():
    """各塔的等级上限（图册与游戏配置一致）。"""
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    fallback = {"哨塔": 100, "散弹塔": 130, "炮塔": 110, "波动塔": 100, "镶嵌塔": 60}
    if not os.path.exists(src):
        return fallback
    cfg = json.load(open(src, encoding="utf-8"))
    out = {}
    for _, name, _, _ in TOWER_DEFS:
        for key, tower in cfg["building"]["tower"].items():
            if tower.get("name") == name:
                out[name] = int(tower.get("lev_max") or fallback[name])
    return out or fallback


def read_tower_rate():
    """塔的基础攻速（配置里的 rate，面板显示值 = rate×10，攻击间隔 = 2500/(rate×10) 帧）。"""
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    fallback = {"哨塔": 18, "散弹塔": 12, "炮塔": 6, "波动塔": 2, "镶嵌塔": 4}
    if not os.path.exists(src):
        return fallback
    cfg = json.load(open(src, encoding="utf-8"))
    out = {}
    for _, name, _, _ in TOWER_DEFS:
        for key, tower in cfg["building"]["tower"].items():
            if tower.get("name") == name and tower.get("rate"):
                out[name] = tower["rate"]
    return out or fallback


def read_wolves():
    """每张图会出哪些狼、狼的血量系数、波难度系数。

    数据来自游戏配置 sys_config.json：
      * umaps[mid].wolf_proportion = [[累计概率, 狼id], ...]（本图只会出这些狼）
      * umaps[mid].pop_max 是这一波的总 pop 预算，随机生成波时按它抽满
      * umaps[mid].yield_val / hard_ness 用于算等级 L = sqrt(ha + hb*进度)
      * wolfs[wid].hp_factor = {a,b,c}，最大血量 = floor((a + L*(b + L*c)) * 难度系数)
    """
    import json

    src = os.path.join(PROJ, "tdsheepvillage-view-master", "src", "assets", "sys_config.json")
    if not os.path.exists(src):
        return {}, {}, {}
    cfg = json.load(open(src, encoding="utf-8"))

    def nm(s):
        return s.split("^")[-1] if isinstance(s, str) else s

    map_wolf = {}
    used = set()
    for mid, m in cfg["umaps"].items():
        prop = [[p, wid] for p, wid in (m.get("wolf_proportion") or [])]
        boss = list(m.get("boss") or [])
        rnd = [[e[0], e[1], e[2]] for e in (m.get("random_boss") or [])]
        map_wolf[nm(m.get("name"))] = {
            "gameId": mid,
            "popMax": m.get("pop_max", 10),
            "hardA": m.get("yield_val", 0),
            "hardB": m.get("hard_ness", 0),
            "scoreMax": m.get("pass_score", 0),
            "prop": prop,
            "boss": boss,
            "randomBoss": rnd,
        }
        for _, wid in prop:
            used.add(wid)
        used.update(boss)
        for e in rnd:
            used.add(e[1])

    wolfs = {}
    for wid in sorted(used):
        w = cfg["wolfs"].get(wid)
        if not w:
            continue
        f = w.get("hp_factor") or {}
        wolfs[wid] = {
            "name": nm(w.get("name")),
            "a": f.get("a", 0), "b": f.get("b", 0), "c": f.get("c", 0),
            "pop": w.get("pop", 1),
            "speed": w.get("speed"),
            "defense": w.get("defense"),
            # 体型：宽 × 高 ≥ 487500 算「沉重」，不受弹簧影响；轻的会被弹簧弹到落点
            "w": w.get("width"), "h": w.get("height"),
            # 带 Fly_* 技能的狼会飞：炮塔、波动塔打不到；散弹塔打它伤害与效果翻倍
            "fly": any(str(s.get("skid", "")).startswith("Fly") for s in (w.get("skills") or [])),
        }
    wave_diff = {k: nm(v.get("desc")) for k, v in (cfg.get("wolf_hard_ness") or {}).items()}
    return map_wolf, wolfs, wave_diff


def export_images(anchors, maps, name_rows):
    if os.path.isdir(MAPS_DIR):
        for f in os.listdir(MAPS_DIR):
            os.remove(os.path.join(MAPS_DIR, f))
    os.makedirs(MAPS_DIR, exist_ok=True)
    pool = sorted([a for a in anchors if a["col"] <= 4], key=lambda a: a["row"])
    for m in maps:
        rows = name_rows.get(m["name"], [])
        best = None
        for idx, a in enumerate(pool):
            if a.get("used"):
                continue
            d = min([abs(a["row"] - r) for r in rows] or [999])
            if best is None or d < best[0]:
                best = (d, idx, a)
        if best is None or best[0] > 3:
            m["img"] = None
            continue
        anchor = best[2]
        anchor["used"] = True
        im = Image.open(os.path.join(MEDIA_DIR, anchor["media"][0]))
        if im.mode in ("RGBA", "LA"):
            box = im.getchannel("A").getbbox()
            if box:
                im = im.crop(box)
        if im.width > 1000:
            im = im.resize((1000, round(im.height * 1000 / im.width)), Image.LANCZOS)
        name = "map_%02d.png" % m["order"]
        im.convert("RGBA").save(os.path.join(MAPS_DIR, name))
        m["img"] = "maps/" + name


def export_references(maps):
    from pypdf import PdfReader

    if not os.path.exists(PDF):
        return
    if os.path.isdir(REFS_DIR):
        for f in os.listdir(REFS_DIR):
            os.remove(os.path.join(REFS_DIR, f))
    os.makedirs(REFS_DIR, exist_ok=True)
    reader = PdfReader(PDF)
    by_no = {m["no"]: m for m in maps}
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            continue
        flat = re.sub(r"\s+", "", text)
        if not flat:
            continue
        # 页面标题格式：第一大图3卡西村… / 12埃特纳村… ，取开头的地图编号
        tail = re.sub(r"^.*?第[一二三四五六七八九十]+大图", "", flat)
        num = re.match(r"(\d{1,2})", tail)
        best = by_no.get(int(num.group(1))) if num else None
        if best is None:
            pos = 10 ** 9
            for m in maps:
                for name in m.get("alias", [m["name"]]):
                    i = flat.find(name)
                    if 0 <= i < pos:
                        best, pos = m, i
        if best is None:
            continue
        if best is None:
            continue
        if best.get("income") is None:
            inc = re.search(r"参考银币收入：([0-9]+|无限)", flat)
            if inc:
                best["income"] = inc.group(1)
        try:
            images = list(page.images)
        except Exception:
            continue
        for i, img in enumerate(images):
            try:
                im = Image.open(BytesIO(img.data)).convert("RGB")
            except Exception:
                continue
            if im.width < 400:
                continue
            if im.width > 1050:
                im = im.resize((1050, round(im.height * 1050 / im.width)), Image.LANCZOS)
            ref_name = "ref_%02d_%d.jpg" % (best["order"], i)
            im.save(os.path.join(REFS_DIR, ref_name), quality=78, optimize=True)
            best["refs"].append("refs/" + ref_name)


def main():
    unzip_media()
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["前线、防线地图"]

    name_rows = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or str(cell.value).startswith("="):
                continue
            name = str(cell.value).strip()
            name = re.split(r"[\uff08(]", name)[0].strip()
            if not name or len(name) > 20:
                continue
            if re.search(r"(村|镇|城|岛|厂|间|林|防线|沼泽|矿山|地带|熔炉|之巅|冰川|谷|桥|冰场|雪原|瀑布)$", name):
                name_rows.setdefault(name, []).append(cell.row)

    blocks = find_blocks(ws)
    candidates = {}
    for block in sorted(blocks, key=lambda b: (b["top"], b["left"])):
        if block["top"] <= 9:
            continue
        raw = read_cell(ws.cell(row=block["top"], column=block["left"]))["text"]
        if not raw:
            continue
        title = re.split(r"[\uff08(]", raw)[0].strip()
        if title not in name_rows:
            continue
        parsed = parse_map(ws, block)
        if parsed is None:
            continue
        candidates.setdefault(title, []).append((block, parsed, raw))

    def version_score(raw):
        """同一张图有旧版/新版两份时，优先用新版。"""
        if "此图为新版" in raw or "以本图为准" in raw:
            return 0
        if "本图为旧版" in raw or "本图有误" in raw:
            return 2
        return 1

    parsed_maps = []
    for title, lst in candidates.items():
        lst.sort(key=lambda x: version_score(x[2]))
        block, parsed, _raw = lst[0]
        parsed_maps.append({
            "order": block["top"],
            "name": title,
            "section": "defend" if block["top"] >= 584 else "front",
            "img": None,
            "refs": [],
            "feat": {},
            "marks": {},
            "spring": {},
            "belt": {},
            "teleColor": {},
            "income": None,
            **parsed,
        })

    # 按《全地图布局参考》的编号顺序排列，并标出所属大图
    by_name = {m["name"]: m for m in parsed_maps}
    if "防线" in by_name:
        by_name["防线"]["alias"] = ["防线（本图即大本营防线）"]
    ordered = []
    for i, name in enumerate(PDF_ORDER):
        m = by_name.get(name)
        if m is None:
            print("  ! 没有解析到地图：", name)
            continue
        m["no"] = i + 1
        m["chapter"] = 1 if i < 11 else (2 if i < 25 else (3 if i < 37 else 4))
        # 村口规则：拉帕斯村两个村口都不能堵，其它图至少留 1 个
        m["villageRule"] = "all" if name == "\u62c9\u5e15\u65af\u6751" else "one"
        m["order"] = i
        ordered.append(m)
    for name in EXTRA_MAPS:
        m = by_name.get(name)
        if m is None:
            print("  ! 没有解析到额外地图：", name)
            continue
        m["no"] = None
        m["chapter"] = 0            # 单独一组，与四个大图并列
        m["order"] = len(ordered)
        ordered.append(m)
    maps = ordered

    anchors = read_drawing_anchors()
    export_images(anchors, maps, name_rows)
    # 布局参考图（refs/）在工具里用不上，分享包里也不需要，所以不再导出。
    # 需要的话把下面这行取消注释即可重新生成。
    # export_references(maps)
    # 大本营防线与 38 号「防线」是同一张图，参考收入沿用
    twin = next((m for m in maps if m["name"] == "防线"), None)
    extra = next((m for m in maps if m["name"] == "大本营防线"), None)
    if twin and extra and not extra.get("income"):
        extra["income"] = twin.get("income")

    # 每张图的出狼配置（大本营防线用「防线」的）
    map_wolf, wolfs, wave_diff = read_wolves()
    for m in maps:
        cfg = map_wolf.get(m["name"])
        if cfg is None and m["name"] == "大本营防线":
            cfg = map_wolf.get("防线")
        m["wolf"] = cfg
        # 传送带方向（人工核对表）
        rule = BELT_DIR.get(m["name"])
        if rule:
            belt = {}
            for pos, fname in (m.get("feat") or {}).items():
                if fname != "传送带":
                    continue
                d = rule if isinstance(rule, str) else rule.get(pos)
                if d:
                    belt[pos] = d
            m["belt"] = belt
        # 个别格子的地形修正
        for pos, t in (CELL_FIX.get(m["name"]) or {}).items():
            rr, cc = [int(x) for x in pos.split("_")]
            row = m["grid"][rr]
            if cc < len(row):
                m["grid"][rr] = row[:cc] + str(t) + row[cc+1:]
                m.get("feat", {}).pop(pos, None)

    front = read_table(wb, "各种塔造价、攻击力", TOWER_DEFS)
    defend = read_table(wb, "防线塔造价、经验", TOWER_DEFS,
                        exp_cols={"哨塔": 3, "散弹塔": 5, "炮塔": 7, "波动塔": 9, "镶嵌塔": 11})

    data = {
        "version": "V1.82",
        "towers": [{"id": k, "name": n} for k, n, _, _ in TOWER_DEFS],
        "typeName": TYPE_NAME,
        "frontCost": {n: front[n]["cost"] for _, n, _, _ in TOWER_DEFS},
        "frontAtk": {n: front[n]["atk"] for _, n, _, _ in TOWER_DEFS},
        "defendCost": {n: defend[n]["cost"] for _, n, _, _ in TOWER_DEFS},
        "defendExp": {n: defend[n]["exp"] for _, n, _, _ in TOWER_DEFS},
        "upgradeTime": read_upgrade_time(wb),
        "towerMax": read_tower_max(),
        "towerRate": read_tower_rate(),
        "gems": read_gems(),
        "baseRange": BASE_RANGE,
        "crossRange": CROSS_RANGE,
        "wolfs": wolfs,
        "wolfSkill": read_wolf_skills(),
        "wolfSkills": read_wolf_skill_list(),
        "waveDiff": wave_diff,
        "gemEffect": read_gem_effects(wb),
        "special": read_special(wb),
        "maps": maps,
    }

    with open(os.path.join(ROOT, "data.js"), "w", encoding="utf-8") as f:
        f.write("// 由 _build/extract_data.py 从《保卫羊村百科全书V1.82.xlsx》自动生成\n")
        f.write("// 如需更新数据，请重新运行 _build/extract_data.py\n")
        f.write("window.YC_DATA = ")
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")

    print("地图 %d 张" % len(maps))
    for m in maps:
        kinds = Counter("".join(m["grid"]))
        print("  %-24s %-6s %2dx%-2d 可走%3d 标记%d 边%d 参考%d %s" % (
            m["name"], m["section"], m["cols"], m["rows"],
            kinds.get("1", 0) + kinds.get("7", 0),
            len(m["labels"]), len(m["edges"]), len(m["refs"]), m["img"]))


if __name__ == "__main__":
    main()
